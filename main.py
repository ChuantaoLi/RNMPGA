"""
模型运行入口

用法:
    python main.py --config config.yaml
    python main.py --config config.yaml --resume   # 检查结果日志, 跳过已有结果
"""

import argparse
import contextlib
import hashlib
import json
import os
import time
import traceback
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from sklearn.model_selection import train_test_split
from tqdm import tqdm

from Model import get_model
from Model.functionHO import build_classifier
from metrics import (
    feature_ratio,
    hypervolume_2d,
    igd,
    mce_metric,
    non_dominated_indices,
)

ROOT = Path(__file__).resolve().parent
HV_REF = (1.0, 1.0)

RESULT_COLUMNS = [
    "Timestamp",
    "Dataset",
    "Model",
    "Params",
    "MCE",
    "FN",
    "HV",
    "IGD",
    "Features_Selected",
    "Features_Total",
    "Runtime_s",
    "MCE_std",
    "FN_std",
    "HV_std",
    "IGD_std",
    "Features_Selected_std",
    "Runtime_s_std",
]


def mean_std(values, mean_digits=6, std_digits=6) -> tuple[float, float]:
    """
    repeats 次独立运行的均值与样本标准差
    """
    vals = np.asarray(values, dtype=float)
    m = float(vals.mean())
    s = float(vals.std(ddof=1)) if vals.size > 1 else 0.0
    return round(m, mean_digits), round(s, std_digits)


def load_config(path: Path) -> dict:
    """
    读取 yaml 配置
    """
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def params_str(params: dict) -> str:
    """
    模型参数序列化字符串
    """
    return json.dumps(params, ensure_ascii=False)


def _checkpoint_path(
    checkpoint_dir: Path, ds_name: str, model_name: str, params: dict
) -> Path:
    digest = hashlib.md5(params_str(params).encode("utf-8")).hexdigest()[:8]
    name = f"checkpoint_{_safe_name(ds_name)}_{_safe_name(model_name)}_{digest}.json"
    return checkpoint_dir / name


def load_repeat_checkpoint(
    checkpoint_dir: Path | None, ds_name: str, model_name: str, params: dict
) -> dict[int, dict]:
    if checkpoint_dir is None:
        return {}
    path = _checkpoint_path(checkpoint_dir, ds_name, model_name, params)
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        return {int(k): v for k, v in payload.get("repeats", {}).items()}
    except Exception as e:
        print(f"  [警告] 无法读取 repeat checkpoint {path}: {e}, 将重新运行缺失 repeat")
        return {}


def save_repeat_checkpoint(
    checkpoint_dir: Path,
    ds_name: str,
    model_name: str,
    params: dict,
    repeats_data: dict[int, dict],
) -> None:
    checkpoint_dir.mkdir(parents=True, exist_ok=True)
    path = _checkpoint_path(checkpoint_dir, ds_name, model_name, params)
    payload = {
        "dataset": ds_name,
        "model": model_name,
        "params": params_str(params),
        "repeats": repeats_data,
    }
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, path)


def _safe_name(s: str) -> str:
    """
    文件名安全化: 非字母数字/_./- 的字符替换为 _
    """
    return "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in s)


def write_curves(
    curves_dir: Path | None,
    ds_name: str,
    model_name: str,
    params: dict,
    repeats: int,
    curves: list,
) -> None:
    """
    导出每 repeat 的逐代适应度曲线 (JSON)。
    """
    if curves_dir is None:
        return
    curves_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.md5(params_str(params).encode("utf-8")).hexdigest()[:8]
    fname = f"curves_{_safe_name(ds_name)}_{_safe_name(model_name)}_{digest}.json"
    payload = {
        "dataset": ds_name,
        "model": model_name,
        "params": params_str(params),
        "repeats": repeats,
        "maxgen": len(curves[0]) if curves and curves[0] else None,
        "curves": curves,
        "note": "curves[i] 为第 i 次重复 (repeat) 的逐代最优适应度 (模型返回的 "
        "'c' 收敛曲线; 多数模型为分类准确率越大越好, 部分为误差/适应度越小越好)",
    }
    tmp = curves_dir / (fname + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, curves_dir / fname)


def write_objective_space(
    obj_dir: Path | None,
    ds_name: str,
    model_name: str,
    params: dict,
    repeats: int,
    D: int,
    run_obj: list,
    obj: np.ndarray,
) -> None:
    """
    导出目标空间数据
    """
    if obj_dir is None:
        return
    obj_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.md5(params_str(params).encode("utf-8")).hexdigest()[:8]
    fname = f"obj_{_safe_name(ds_name)}_{_safe_name(model_name)}_{digest}.json"
    payload = {
        "dataset": ds_name,
        "model": model_name,
        "params": params_str(params),
        "repeats": repeats,
        "D": int(D),
        "objectives": ["|s|/D", "1-Acc"],  # 论文式 (1): 特征占比, 分类错误率
        "obj": [[float(a), float(b)] for a, b in obj],
        "fronts": [[[float(a), float(b)] for a, b in f] for f in run_obj],
        "note": "目标空间坐标 (|s|/D, 1-Acc); obj 为合并全部 repeats 的非支配解"
        "(该模型自身的近似 PF), fronts[i] 为第 i 次重复的前沿",
    }
    tmp = obj_dir / (fname + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, obj_dir / fname)


def resolve_datasets(cfg: dict) -> list[str]:
    """
    解析配置中的数据集列表; 填 auto 或留空时运行 Dataset 下全部 csv
    """
    names = cfg.get("datasets", [])
    if isinstance(names, str):
        names = [names]
    if not names or any(str(n).strip().lower() == "auto" for n in names):
        names = sorted(p.name for p in (ROOT / "Dataset").glob("*.csv"))
    return names


def load_dataset(name: str) -> tuple[np.ndarray, np.ndarray]:
    """
    加载数据集: 特征取除最后一列, 标签取最后一列 (兼容不同列名)
    """
    path = ROOT / "Dataset" / name
    if not path.exists():
        raise FileNotFoundError(f"数据集不存在: {path}")
    df = pd.read_csv(path)
    feat = df.iloc[:, :-1].to_numpy(dtype=float)
    label = df.iloc[:, -1].to_numpy().ravel()
    if np.isnan(feat).any():
        raise ValueError(f"{name}: 特征矩阵含有 NaN 值")
    return feat, label


def run_one(
    ds_name: str,
    model_name: str,
    model_fn,
    params: dict,
    split_cfg: dict,
    repeats: int,
    show_bar: bool,
    curves_dir: Path | None = None,
    obj_dir: Path | None = None,
    repeat_checkpoint: dict[int, dict] | None = None,
    checkpoint_dir: Path | None = None,
) -> dict:
    """
    对单个 (数据集, 模型) 组合: 划分数据 -> 独立运行 repeats 次 -> 论文指标评估
    """
    feat, label = load_dataset(ds_name)
    D = feat.shape[1]

    stratify = split_cfg.get("stratify", True)
    split_kwargs = dict(
        test_size=split_cfg.get("test_size", 0.3),
        random_state=split_cfg.get("random_state", None),
    )
    try:
        xtr, xte, ytr, yte = train_test_split(
            feat, label, stratify=label if stratify else None, **split_kwargs
        )
    except ValueError:
        print(f"  [警告] {ds_name}: 按标签分层失败, 改为随机划分")
        xtr, xte, ytr, yte = train_test_split(feat, label, **split_kwargs)

    cls_opts = dict(params)

    def evaluate(sel_idx):
        if sel_idx.size == 0:
            return 0.0, 1.0  # 未选中任何特征: 准确率 0, MCE 视为 1
        mdl = build_classifier(cls_opts)
        mdl.fit(xtr[:, sel_idx], np.asarray(ytr).ravel())
        y_pred = mdl.predict(xte[:, sel_idx])
        acc = float(np.mean(np.asarray(yte).ravel() == y_pred))
        mce = mce_metric(yte, y_pred)
        return acc, mce

    run_pfs = []
    run_times = []
    run_curves = []
    repeat_checkpoint = repeat_checkpoint or {}
    for rep in range(repeats):
        saved = repeat_checkpoint.get(rep)
        if saved is not None:
            run_pfs.append([np.asarray(p, dtype=int) for p in saved["pf"]])
            run_times.append(float(saved["elapsed"]))
            run_curves.append([float(v) for v in saved.get("curve", [])])
            print(
                f"  [resume] {ds_name} | {model_name} | run {rep + 1}/{repeats} 已完成，跳过"
            )
            continue

        opts = dict(params)
        if opts.get("seed") is not None:

            opts["seed"] = int(opts["seed"]) + rep
        opts["fold"] = {"xt": xtr, "yt": ytr, "xv": xte, "yv": yte}

        bar = None
        if show_bar:
            total = int(opts.get("T", 0) or opts.get("maxGen", 0) or 0)
            if total > 0:
                bar = tqdm(
                    total=total,
                    desc=f"[{ds_name} | {model_name} | run {rep + 1}/{repeats}]",
                    unit="iter",
                    leave=False,
                )
                opts["progress"] = lambda n: bar.update(n - bar.n)

        try:
            t0 = time.perf_counter()
            fmdl = model_fn(feat, label, opts)
            elapsed = time.perf_counter() - t0
        finally:
            if bar is not None:
                bar.close()

        pf = fmdl.get("pf")
        if pf is None:
            sf = np.asarray(fmdl["sf"], dtype=int)
            if sf.size == 0:
                raise RuntimeError(f"{model_name}: 未选中任何特征, 无法评估")
            pf = [sf]
        pf_arrays = [np.asarray(p, dtype=int) for p in pf]
        run_pfs.append(pf_arrays)
        run_times.append(elapsed)
        curve = [float(v) for v in fmdl["c"]] if fmdl.get("c") is not None else []
        run_curves.append(curve)
        if checkpoint_dir is not None:
            repeat_checkpoint[rep] = {
                "pf": [p.tolist() for p in pf_arrays],
                "elapsed": float(elapsed),
                "curve": curve,
            }
            save_repeat_checkpoint(
                checkpoint_dir, ds_name, model_name, params, repeat_checkpoint
            )

    acc_cache: dict[tuple, tuple] = {}
    seen = set()
    for pf in run_pfs:
        for sel in pf:
            seen.add(tuple(int(i) for i in sel))
    for key in seen:
        sel = np.asarray(key, dtype=int)
        acc_cache[key] = evaluate(sel)

    all_subsets = list(seen)
    all_obj = np.array([[len(k) / D, 1.0 - acc_cache[k][0]] for k in all_subsets])
    nd = non_dominated_indices(all_obj)
    true_pf_obj = all_obj[nd]

    mces, fns, hvs, igds, n_sels = [], [], [], [], []
    run_obj = []
    for pf in run_pfs:
        uniq = {}
        for sel in pf:
            key = tuple(int(i) for i in sel)
            uniq[key] = acc_cache[key]
        obj_uniq = np.array([[len(k) / D, 1.0 - uniq[k][0]] for k in uniq])
        run_obj.append(obj_uniq)

        best_key = min(uniq, key=lambda k: uniq[k][1])
        mces.append(uniq[best_key][1])
        fns.append(feature_ratio(len(best_key), D))
        n_sels.append(len(best_key))
        hvs.append(hypervolume_2d(obj_uniq, HV_REF))
        igds.append(igd(obj_uniq, true_pf_obj))

    mce_m, mce_s = mean_std(mces)
    fn_m, fn_s = mean_std(fns)
    hv_m, hv_s = mean_std(hvs)
    igd_m, igd_s = mean_std(igds)
    nsel_m, nsel_s = mean_std(n_sels, 6, 6)
    rt_m, rt_s = mean_std(run_times, 6, 6)

    if curves_dir is not None:
        write_curves(curves_dir, ds_name, model_name, params, repeats, run_curves)

    if obj_dir is not None:
        write_objective_space(
            obj_dir, ds_name, model_name, params, repeats, D, run_obj, true_pf_obj
        )

    return {
        "Dataset": ds_name,
        "Model": model_name,
        "Params": params_str(params),
        "MCE": mce_m,
        "FN": fn_m,
        "HV": hv_m,
        "IGD": igd_m,
        "Features_Selected": nsel_m,
        "Features_Total": int(D),
        "Runtime_s": rt_m,
        "MCE_std": mce_s,
        "FN_std": fn_s,
        "HV_std": hv_s,
        "IGD_std": igd_s,
        "Features_Selected_std": nsel_s,
        "Runtime_s_std": rt_s,
    }


def append_row(path: Path, row: dict) -> None:
    """向 xlsx 追加一行; 文件不存在时创建并写入表头"""
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(RESULT_COLUMNS)
    ws.append([row[c] for c in RESULT_COLUMNS])
    wb.save(path)


def load_existing_keys(path: Path) -> set[tuple[str, str, str]]:
    """
    --resume 用: 读取日志中已有的 (Dataset, Model, Params) 组合集合
    """
    if not path.exists():
        return set()
    try:
        df = pd.read_excel(path)
    except Exception as e:
        print(f"  [警告] 无法读取日志 {path}: {e}, 将全部运行")
        return set()
    if df.empty or not {"Dataset", "Model", "Params"} <= set(df.columns):
        return set()
    keys = set()
    for ds, m, p in zip(df["Dataset"], df["Model"], df["Params"]):
        if pd.isna(ds) or pd.isna(m) or pd.isna(p):
            continue  # 跳过空行
        keys.add((str(ds), str(m), str(p)))
    return keys


def apply_thread_limits(n_jobs):
    """
    把底层数值库 (MKL/OpenBLAS 等) 的线程数限制到 n_jobs
    """
    if n_jobs is None:
        return contextlib.nullcontext()
    for var in (
        "MKL_NUM_THREADS",
        "OMP_NUM_THREADS",
        "OPENBLAS_NUM_THREADS",
        "NUMEXPR_NUM_THREADS",
    ):
        os.environ.setdefault(var, str(n_jobs))
    try:
        from threadpoolctl import threadpool_limits
    except ImportError:
        print("  [提示] 未安装 threadpoolctl, 仅通过环境变量限制线程数")
        return contextlib.nullcontext()
    return threadpool_limits(limits=n_jobs, user_api="blas")


def report_thread_limits() -> None:
    """
    打印限制后底层库 (MKL/OpenBLAS) 的实际线程数, 便于确认生效
    """
    try:
        from threadpoolctl import threadpool_info
    except ImportError:
        return
    for info in threadpool_info():
        if info["user_api"] == "blas":
            print(
                f"  [线程] {info['internal_api']}: {info['num_threads']} "
                f"线程 (由 compute.n_jobs 限制)"
            )


def main() -> None:
    parser = argparse.ArgumentParser(description="特征筛选统一运行入口")
    parser.add_argument("--config", default="config.yaml", help="配置文件路径")
    parser.add_argument(
        "--resume",
        action="store_true",
        help="检查结果日志, 跳过已有结果的 (数据集, 模型, 参数) 组合",
    )
    args = parser.parse_args()

    cfg = load_config(ROOT / args.config)

    n_jobs = (cfg.get("compute") or {}).get("n_jobs")
    if n_jobs is not None:
        n_jobs = int(n_jobs)
        if n_jobs < 1:
            raise ValueError(f"compute.n_jobs 必须为正整数, 得到 {n_jobs}")
    with apply_thread_limits(n_jobs):
        if n_jobs is not None:
            report_thread_limits()
        run_experiments(cfg, args)


def run_experiments(cfg: dict, args: argparse.Namespace) -> None:
    """
    主实验循环: 对每个 (数据集, 模型) 组合运行特征选择并记录结果
    """
    ds_names = resolve_datasets(cfg)
    out_cfg = cfg.get("output", {})
    out_dir = ROOT / out_cfg.get("dir", "Results")
    log_path = out_dir / out_cfg.get("log_file", "results_log.xlsx")
    curves_dir = out_dir / "curves"
    objective_dir = out_dir / "objective_space"
    checkpoint_dir = out_dir / "repeat_checkpoints"
    show_bar = cfg.get("progress_bar", True)
    split_cfg = cfg.get("split", {})
    repeats = int(cfg.get("repeats", 30))  # 论文: 30 次独立运行

    print(
        f"数据集 {len(ds_names)} 个, 每个组合独立运行 {repeats} 次, "
        f"结果日志: {log_path}\n"
    )

    existing = load_existing_keys(log_path) if args.resume else set()
    if args.resume:
        print(
            f"  [resume] 日志中已有 {len(existing)} 条结果, "
            f"匹配 (数据集, 模型, 参数) 的组合将跳过\n"
        )

    total_runs = 0
    skipped = 0
    classifier_cfg = cfg.get("classifier", {})
    classifier_type = classifier_cfg.get("type", "dt")
    classifier_params = classifier_cfg.get("params", {}) or {}
    for model_cfg in cfg.get("models", []):
        model_name = model_cfg["name"]
        try:
            model_fn = get_model(model_name)
        except KeyError as e:
            print(e)
            continue
        params = {
            **model_cfg.get("params", {}),
            **classifier_params,
            "classifier": classifier_type,
        }

        for ds in ds_names:
            if args.resume and (ds, model_name, params_str(params)) in existing:
                print(f"  [跳过] {ds} | {model_name} (结果已存在于日志)")
                skipped += 1
                continue

            ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            try:
                repeat_checkpoint = (
                    load_repeat_checkpoint(checkpoint_dir, ds, model_name, params)
                    if args.resume
                    else {}
                )
                row = run_one(
                    ds,
                    model_name,
                    model_fn,
                    params,
                    split_cfg,
                    repeats,
                    show_bar,
                    curves_dir,
                    objective_dir,
                    repeat_checkpoint,
                    checkpoint_dir,
                )
            except Exception as e:
                print(
                    f"  [失败] {ds} | {model_name}: {type(e).__name__}: {e}", flush=True
                )
                traceback.print_exc()
                continue

            row = {"Timestamp": ts, **row}
            append_row(log_path, row)  # 追加式日志 (续写, 不覆盖)
            total_runs += 1

            print(
                f"  MCE={row['MCE']:.4f}±{row['MCE_std']:.4f} | "
                f"FN={row['FN']:.4f}±{row['FN_std']:.4f} | "
                f"HV={row['HV']:.4f}±{row['HV_std']:.4f} | "
                f"IGD={row['IGD']:.4f}±{row['IGD_std']:.4f} | "
                f"特征 {row['Features_Selected']:.1f}±{row['Features_Selected_std']:.1f}/"
                f"{row['Features_Total']} | {row['Runtime_s']:.1f}±{row['Runtime_s_std']:.1f}s"
            )

    skip_note = f", 跳过 {skipped} 条已有结果" if args.resume else ""
    print(f"\n完成, 共 {total_runs} 条结果{skip_note}")
    print(f"  日志文件: {log_path}")


if __name__ == "__main__":
    main()
